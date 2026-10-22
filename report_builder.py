"""
report_builder.py - turns a messy sales export into a clean excel report

usage:
    python report_builder.py orders_export.csv -o sales_report.xlsx

what it does:
    - fixes mixed date formats, prices like "₪1,200.00" / "1200,5" / "ILS 300"
    - trims and normalises customer and product names
    - drops duplicate orders, moves broken rows to an "Issues" sheet
    - builds Summary, By month, By product sheets with a chart
"""
import argparse
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DATE_FORMATS = ["%d/%m/%Y", "%Y-%m-%d", "%d %b %Y", "%d.%m.%Y"]
HEADER_FILL = PatternFill("solid", fgColor="1D5A45")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def parse_date(value):
    # try every known format, return None if nothing fits
    value = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def parse_amount(value):
    # "₪1,200.50" -> 1200.5, "240,0" -> 240.0, "ILS 300" -> 300.0
    text = re.sub(r"[^\d,.\-]", "", str(value))
    if not text:
        return None
    if "," in text and "." in text:
        text = text.replace(",", "")          # thousands separator
    elif "," in text:
        text = text.replace(",", ".")         # decimal comma
    try:
        return round(float(text), 2)
    except ValueError:
        return None


def clean(df):
    df = df.rename(columns=lambda c: c.strip().lower().replace(" ", "_"))
    df["customer"] = df["customer"].str.strip().str.title()
    df["product"] = df["product"].str.strip().str.capitalize()
    df["qty"] = pd.to_numeric(df["qty"], errors="coerce").fillna(1).astype(int)
    df["date_clean"] = df["date"].map(parse_date)
    df["total_clean"] = df["total"].map(parse_amount)

    problems = []
    bad_date = df["date_clean"].isna()
    bad_total = df["total_clean"].isna()
    for i in df.index[bad_date]:
        problems.append((i, "unreadable date"))
    for i in df.index[bad_total & ~bad_date]:
        problems.append((i, "missing total"))

    issues = df.loc[[i for i, _ in problems], ["order_id", "date", "customer", "product", "total"]].copy()
    issues["problem"] = [p for _, p in problems]

    good = df[~bad_date & ~bad_total].copy()
    before = len(good)
    good = good.drop_duplicates(subset="order_id", keep="first")
    dupes = before - len(good)

    good = good.rename(columns={"date_clean": "order_date", "total_clean": "total_ils"})
    good = good[["order_id", "order_date", "customer", "product", "qty", "total_ils"]]
    good = good.sort_values("order_date").reset_index(drop=True)
    return good, issues, dupes


def build_report(good, issues, dupes, raw_rows, out_path):
    good["month"] = pd.to_datetime(good["order_date"]).dt.to_period("M").astype(str)
    by_month = good.groupby("month").agg(orders=("order_id", "count"), revenue_ils=("total_ils", "sum")).reset_index()
    by_product = (good.groupby("product")
                  .agg(units=("qty", "sum"), revenue_ils=("total_ils", "sum"))
                  .sort_values("revenue_ils", ascending=False).reset_index())

    summary = pd.DataFrame({
        "metric": ["Rows in export", "Clean orders", "Duplicates removed", "Rows with issues",
                   "Total revenue, ₪", "Average order, ₪", "Best month", "Top product"],
        "value": [raw_rows, len(good), dupes, len(issues),
                  round(good["total_ils"].sum(), 2), round(good["total_ils"].mean(), 2),
                  by_month.loc[by_month["revenue_ils"].idxmax(), "month"],
                  by_product.iloc[0]["product"]],
    })

    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        by_month.to_excel(xw, sheet_name="By month", index=False)
        by_product.to_excel(xw, sheet_name="By product", index=False)
        good.drop(columns="month").to_excel(xw, sheet_name="Orders", index=False)
        issues.to_excel(xw, sheet_name="Issues", index=False)

    style_workbook(out_path, len(by_month))


def style_workbook(path, month_rows):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col) + 3
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(width, 40)
        ws.freeze_panes = "A2"
        money_cols = [c.column for c in ws[1] if c.value and ("ils" in str(c.value) or "₪" in str(c.value))]
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.column in money_cols or isinstance(c.value, float):
                    c.number_format = "#,##0.00"

    chart = BarChart()
    chart.title, chart.y_axis.title = "Revenue by month, ₪", "₪"
    ws = wb["By month"]
    chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=month_rows + 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=month_rows + 1))
    chart.width, chart.height, chart.legend = 18, 8, None
    chart.series[0].graphicalProperties.solidFill = "1D5A45"
    ws.add_chart(chart, "E2")
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description="clean a sales export and build an excel report")
    ap.add_argument("csv", type=Path)
    ap.add_argument("-o", "--out", type=Path, default=Path("sales_report.xlsx"))
    args = ap.parse_args()

    raw = pd.read_csv(args.csv, dtype=str)
    good, issues, dupes = clean(raw)
    build_report(good, issues, dupes, len(raw), args.out)

    print(f"read      {len(raw)} rows from {args.csv.name}")
    print(f"removed   {dupes} duplicate orders")
    print(f"flagged   {len(issues)} rows -> 'Issues' sheet")
    print(f"revenue   ₪{good['total_ils'].sum():,.2f} across {len(good)} orders")
    print(f"saved     {args.out}")


if __name__ == "__main__":
    main()
